from __future__ import annotations

import csv
import io
from itertools import chain
from typing import Iterable


FIELDS = ("id", "product", "price", "unit", "currency", "source", "supplier", "source_date", "collected_at", "variation", "variation_pct", "is_unchanged", "notes", "data_kind", "benchmark_code", "instrument")
MAX_EXCEL_DATA_ROWS = 1_048_575  # One header plus data must fit a worksheet.


def csv_bytes(rows: Iterable[dict]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=FIELDS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def xlsx_bytes(rows: Iterable[dict]) -> bytes:
    # Validate/open a lazy query before creating temporary worksheet files.
    iterator = iter(rows)
    first = next(iterator, None)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.cell import WriteOnlyCell
    except ImportError as exc:
        raise RuntimeError("Installez openpyxl pour activer l'export Excel (.xlsx).") from exc
    output = io.BytesIO()
    workbook = Workbook(write_only=True)
    headers = ["ID", "Produit", "Prix / indice", "Unité", "Devise", "Source", "Fournisseur", "Date source", "Date collecte", "Variation", "Variation %", "Inchangée", "Remarque", "Type de donnée", "Code indicateur", "Instrument"]
    def new_sheet():
        number = len(workbook.worksheets) + 1
        sheet = workbook.create_sheet("Relevés" if number == 1 else f"Relevés {number}")
        sheet.freeze_panes = "A2"
        for column, width in {"A": 8, "B": 16, "C": 14, "D": 16, "F": 34, "G": 24, "H": 16, "I": 24, "M": 46, "P": 28}.items():
            sheet.column_dimensions[column].width = width
        cells = [WriteOnlyCell(sheet, value=header) for header in headers]
        for cell in cells:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B2E4F")
        sheet.append(cells)
        return sheet

    sheet = new_sheet()
    data_rows = 0
    for row in chain(() if first is None else (first,), iterator):
        if data_rows == MAX_EXCEL_DATA_ROWS:
            sheet = new_sheet()
            data_rows = 0
        sheet.append([row.get(field, "") for field in FIELDS])
        data_rows += 1
    workbook.save(output)
    return output.getvalue()
