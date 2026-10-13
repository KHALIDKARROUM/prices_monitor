from datetime import datetime, timedelta, timezone
import csv
import http.client
import io
import threading
from price_monitor.http_wsgi import WSGIServer
from urllib.parse import urlencode

import pytest
from openpyxl import load_workbook

from price_monitor.services.forecasting import build_forecast
from price_monitor.services import export


def insert_many(database, rows):
    connection = database.connect()
    try:
        connection.executemany(
            """INSERT INTO price_observations
            (product,instrument,price,unit,source,supplier,source_date,collected_at)
            VALUES(?,?,?,?,?,?,?,?)""", rows,
        )
        connection.commit()
    finally:
        connection.close()


def reading(price, day, second=0, source="History limit fixture"):
    return ("brent", "brent_spot", price, "USD/baril", source, "Target",
            day.date().isoformat(), (day + timedelta(seconds=second)).isoformat())


def test_forecast_aggregates_dense_intraday_history_before_applying_limit(empty_database):
    now = datetime.now(timezone.utc).replace(microsecond=0)
    old = now - timedelta(days=50)
    rows = [reading(80 + i / 10000, old, i) for i in range(5001)]
    rows += [reading(90 + i / 100, now - timedelta(days=39 - i)) for i in range(40)]
    insert_many(empty_database, rows)

    result = build_forecast(empty_database)["products"]["brent"]
    assert result["status"] == "ready"
    assert result["history_count"] == 41
    assert result["forecast_origin"] == now.date().isoformat()
    assert result["last_actual"]["value"] == 90.39
    assert result["history"][0]["value"] == 80.5

    recent = empty_database.observations("brent", days=90, limit=2, instrument="brent_spot", daily=True)
    assert [row["source_date"] for row in recent] == [(now - timedelta(days=1)).date().isoformat(), now.date().isoformat()]


def test_raw_history_limits_keep_newest_records_and_return_chronological_order(empty_database):
    now = datetime.now(timezone.utc).replace(microsecond=0)
    insert_many(empty_database, [reading(80 + i, now - timedelta(days=3 - i)) for i in range(4)])
    rows = empty_database.observations("brent", days=30, limit=2)
    assert [row["price"] for row in rows] == [82, 83]
    assert len(empty_database.observations("brent", days=30, limit=None)) == 4


def test_daily_aggregation_keeps_provider_priority_and_latest_intraday_receipt(empty_database):
    now = datetime.now(timezone.utc).replace(microsecond=0)
    insert_many(empty_database, [
        reading(80, now, 0, "EIA/FRED - DCOILBRENTEU"),
        reading(81, now, 1, "EIA/FRED - DCOILBRENTEU"),
        reading(900, now, 2, "U.S. EIA - PET.RBRTE.D"),
    ])
    rows = empty_database.observations("brent", days=30, limit=1, daily=True)
    assert len(rows) == 1
    assert rows[0]["price"] == 81
    assert rows[0]["source"] == "EIA/FRED - DCOILBRENTEU"


@pytest.fixture
def large_history(empty_database):
    now = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    rows = [reading(80 + i / 10000, now, i) for i in range(10003)]
    rows.append(reading(10, now - timedelta(days=1)))
    rows.append(("gasoil", "diesel_spot", 2, "USD/gallon", "Excluded", "Other", now.date().isoformat(), now.isoformat()))
    insert_many(empty_database, rows)
    return empty_database, now


def test_full_history_iteration_has_no_page_cutoff_and_keeps_filters(large_history):
    db, now = large_history
    filters = {"product": "brent", "supplier": "Target", "source": "History limit", "instrument": "brent_spot",
               "data_kind": "quote", "date_from": (now - timedelta(days=1)).date().isoformat(), "date_to": now.date().isoformat(),
               "min_price": 10, "max_price": 100}
    rows = list(db.iter_history(**filters))
    assert len(rows) == 10004
    assert len({row["id"] for row in rows}) == 10004
    assert rows[-1]["price"] == 10
    assert db.history(**filters, page_size=10000)["total"] == 10004
    with pytest.raises(ValueError, match="Instrument"):
        list(db.iter_history(product="brent", instrument="diesel_spot"))


def test_period_comparison_includes_dates_beyond_the_old_ten_thousand_row_cutoff(large_history):
    db, now = large_history
    start, end = (now - timedelta(days=1)).date().isoformat(), now.date().isoformat()
    result = db.compare_periods(start, end, start, end, instrument="brent_spot")["products"]["brent"]
    assert result["period_a"]["count"] == 2
    assert result["period_a"]["average"] == 45.5001
    assert result["average_change"] == 0


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_download_exports_all_filtered_rows_above_ten_thousand(large_history, monkeypatch, extension):
    from price_monitor import app
    db, now = large_history
    monkeypatch.setattr(app, "database", db)
    server = WSGIServer(("127.0.0.1", 0), app.Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    client = http.client.HTTPConnection(*server.server_address, timeout=30)
    try:
        query = urlencode({"product": "brent", "supplier": "Target", "source": "History limit", "instrument": "brent_spot",
                           "data_kind": "quote", "date_from": (now - timedelta(days=1)).date().isoformat(), "date_to": now.date().isoformat(),
                           "min_price": 10, "max_price": 100})
        client.request("GET", f"/export.{extension}?{query}")
        response = client.getresponse()
        body = response.read()
        assert response.status == 200
        assert int(response.getheader("Content-Length")) == len(body)
        if extension == "csv":
            rows = list(csv.DictReader(io.StringIO(body.decode("utf-8-sig"))))
            ids = [int(row["id"]) for row in rows]
            assert rows[-1]["price"] == "10.0"
            assert {row["instrument"] for row in rows} == {"brent_spot"}
        else:
            workbook = load_workbook(io.BytesIO(body), read_only=True)
            try:
                rows = list(workbook["Relevés"].iter_rows(min_row=2, values_only=True))
                ids = [row[0] for row in rows]
                assert rows[-1][2] == 10
                assert {row[-1] for row in rows} == {"brent_spot"}
            finally:
                workbook.close()
        assert len(ids) == len(set(ids)) == 10004
    finally:
        client.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_excel_rolls_over_without_losing_records_or_headers(monkeypatch):
    monkeypatch.setattr(export, "MAX_EXCEL_DATA_ROWS", 2)
    rows = ({"id": i, "product": "brent", "price": 80 + i, "instrument": "brent_spot"} for i in range(5))
    workbook = load_workbook(io.BytesIO(export.xlsx_bytes(rows)))
    assert workbook.sheetnames == ["Relevés", "Relevés 2", "Relevés 3"]
    actual = []
    for sheet in workbook:
        assert sheet.freeze_panes == "A2"
        assert sheet["A1"].value == "ID"
        assert sheet["A1"].font.bold is True
        actual.extend(row[0] for row in sheet.iter_rows(min_row=2, values_only=True))
    assert actual == list(range(5))


def test_history_iterator_closes_the_read_connection_when_cancelled(empty_database, monkeypatch):
    import sqlite3
    now = datetime.now(timezone.utc)
    insert_many(empty_database, [reading(80, now)])
    connection = empty_database.connect()
    monkeypatch.setattr(empty_database, "connect", lambda: connection)
    rows = empty_database.iter_history()
    assert next(rows)["price"] == 80
    rows.close()
    with pytest.raises(sqlite3.ProgrammingError, match="closed"):
        connection.execute("SELECT 1")


def test_mysql_export_uses_unbuffered_results_and_closes_a_cancelled_read():
    from price_monitor.database import MySQLDatabase, _MySQLConnection

    class Cursor:
        def execute(self, statement, params):
            self.statement, self.params = statement, params
        def fetchmany(self, size):
            return [{"id": 1, "product": "brent"}, {"id": 2, "product": "brent"}]

    class NativeConnection:
        def __init__(self):
            self.reader, self.closed = Cursor(), False
        def cursor(self, **kwargs):
            self.options = kwargs
            return self.reader
        def close(self):
            self.closed = True

    native = NativeConnection()
    db = object.__new__(MySQLDatabase)
    db.connect = lambda: _MySQLConnection(native)
    rows = db.iter_history(product="brent", instrument="brent_spot")
    assert next(rows)["id"] == 1
    assert native.options == {"dictionary": True, "buffered": False}
    assert native.reader.params == ("brent", "brent_spot")
    assert "product=%s" in native.reader.statement
    assert "LIMIT" not in native.reader.statement
    rows.close()
    assert native.closed is True
