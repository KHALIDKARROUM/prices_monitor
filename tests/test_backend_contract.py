"""Run the same business regressions on SQLite and a real isolated MySQL DB."""
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from datetime import datetime, timedelta, timezone
import base64
import csv
import io
import os
from uuid import uuid4

from argon2 import PasswordHasher
from openpyxl import load_workbook
import pytest

from price_monitor import access_store, database as store
from price_monitor.config import MYSQL_SETTINGS
from price_monitor.database import SQLiteDatabase, MySQLDatabase, _get_mysql_connector
from price_monitor.quotes import QuoteConflict
from price_monitor.services.alerts import AlertManager
from price_monitor.services.forecasting import build_forecast
from price_monitor.services.export import csv_bytes, xlsx_bytes

@pytest.fixture(params=['sqlite', pytest.param('mysql', marks=pytest.mark.mysql)])
def contract_db(request, tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'REAL_DATA_PATH', tmp_path / 'absent.csv')
    monkeypatch.setattr(store, 'BITUMEN_DATA_PATH', tmp_path / 'absent-bitumen.csv')
    monkeypatch.setattr(store, 'DEMO_MODE', False)
    monkeypatch.setattr(access_store, 'PASSWORDS', PasswordHasher(time_cost=1, memory_cost=8192, parallelism=1))
    if request.param == 'sqlite':
        yield SQLiteDatabase(tmp_path / 'contract.db')
        return
    if os.getenv('PRICE_MONITOR_TEST_MYSQL') != '1':
        pytest.skip('Set PRICE_MONITOR_TEST_MYSQL=1 for isolated real MySQL regressions')
    # Never truncate a configured database: create and drop only our fresh name.
    name = 'pm_test_' + uuid4().hex
    settings = replace(MYSQL_SETTINGS, database=name)
    try:
        yield MySQLDatabase(settings)
    finally:
        assert name.startswith('pm_test_') and len(name) == 40 and name[8:].isalnum()
        conn = _get_mysql_connector().connect(**settings.connector_kwargs(include_database=False))
        try:
            conn.cursor().execute(f'DROP DATABASE IF EXISTS `{name}`')
        finally:
            conn.close()

def observation(instrument, price, at, second=0):
    sources = {'brent_spot': 'EIA/FRED - DCOILBRENTEU', 'brent_futures': 'Yahoo Finance - Brent futures BZ=F'}
    return {'product': 'brent', 'instrument': instrument, 'source': sources[instrument], 'price': price,
            'source_date': at.date().isoformat(), 'collected_at': (at + timedelta(seconds=second)).isoformat()}

def commercial_quote(**changes):
    day = datetime.now(timezone.utc).date()
    return {'product': 'bitume', 'supplier': 'Contract supplier', 'reference': 'Regression', 'grade': '50/70',
        'specification': 'EN 12591', 'origin': 'Maroc', 'incoterm': 'DAP', 'incoterm_place': 'Casablanca',
        'destination': 'Casablanca', 'payment_terms': '30 jours', 'packaging': 'bulk', 'document_type': 'quotation',
        'quote_date': day.isoformat(), 'valid_until': (day + timedelta(days=30)).isoformat(),
        'delivery_date': (day + timedelta(days=7)).isoformat(), 'quantity': 10, 'unit_price': 500,
        'currency': 'USD', 'exchange_rate': 1, **changes}

def test_instrument_switch_and_unchanged_thresholds_do_not_fabricate_events(contract_db):
    day = datetime.now(timezone.utc)
    for instrument in ('brent_spot', 'brent_futures'):
        contract_db.create_alert_rule({'product': 'brent', 'instrument': instrument, 'direction': 'above', 'threshold': 85})
    manager = AlertManager(contract_db)
    rows = [observation('brent_spot', 80, day), observation('brent_futures', 90, day, 1)]
    stored = [contract_db.insert_observation(row) for row in rows]
    assert stored[1]['variation_pct'] is None  # formerly a false +12.5%
    assert [(e.instrument, e.status) for e in manager.evaluate_prices(rows, dispatch=False)] == [('brent_futures', 'triggered')]
    assert manager.evaluate_prices(list(reversed(rows)), dispatch=False) == []
    repeated = contract_db.insert_observation(observation('brent_futures', 90, day, 2))
    assert repeated['id'] == stored[1]['id'] and repeated['variation_pct'] is None
    assert contract_db.dashboard(instruments={'brent': 'brent_spot'})['metrics']['brent']['current'] == 80

def test_dense_history_forecast_and_exports_keep_the_latest_and_all_requested_rows(contract_db):
    day = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    old = day - timedelta(days=50)
    rows = [('brent', 'brent_spot', 80 + i / 10000, 'USD/baril', 'Regression source', 'Private supplier',
             old.date().isoformat(), contract_db._datetime_value((old + timedelta(seconds=i)).isoformat())) for i in range(10001)]
    rows += [('brent', 'brent_spot', 90 + i / 100, 'USD/baril', 'Regression source', 'Private supplier',
              (day - timedelta(days=39-i)).date().isoformat(), contract_db._datetime_value((day - timedelta(days=39-i)).isoformat())) for i in range(40)]
    with contract_db.connect() as conn:
        conn.executemany('INSERT INTO price_observations(product,instrument,price,unit,source,supplier,source_date,collected_at) VALUES(?,?,?,?,?,?,?,?)', rows)
    forecast = build_forecast(contract_db)['products']['brent']
    assert forecast['history_count'] == 41
    assert forecast['forecast_origin'] == day.date().isoformat()
    assert forecast['last_actual']['value'] == 90.39
    assert all(s['selected_model'] == 'naive' for s in forecast['model_selection'].values())
    assert any(w['code'] == 'uncalibrated_intervals' for s in forecast['model_selection'].values() for w in s['warnings'])
    csv_rows = list(csv.DictReader(io.StringIO(csv_bytes(contract_db.iter_history(supplier='Private supplier')).decode('utf-8-sig'))))
    assert len(csv_rows) == 10041 and csv_rows[0]['source_date'] == day.date().isoformat()
    workbook = load_workbook(io.BytesIO(xlsx_bytes(contract_db.iter_history(supplier='Private supplier'))), read_only=True)
    try:
        assert sum(1 for _ in workbook.active.iter_rows()) == 10042
    finally:
        workbook.close()

def test_versioned_quotes_documents_comparisons_and_complete_budget(contract_db):
    quote = contract_db.create_quote(commercial_quote(transport_cost=150, insurance_cost=25, inspection_cost=50,
        duties_cost=100, taxes_cost=20, storage_cost=30, heating_cost=40, finance_cost=5, other_costs=[{'label': 'Port', 'amount': 80}]))
    assert quote['total_cost_usd'] == 5500
    second = contract_db.create_quote(commercial_quote(supplier='Alternative', unit_price=490))
    assert contract_db.compare_quotes([quote['id'], second['id']])['equivalent']
    document = contract_db.attach_quote_document(quote['id'], {'expected_version': 1, 'filename': 'quote.txt',
        'document_type': 'quotation', 'content_base64': base64.b64encode(b'Private supporting quote').decode(), 'change_reason': 'Supporting evidence'})
    assert contract_db.quote_document(quote['id'], document['documents'][0]['id'])['content'] == b'Private supporting quote'
    revised = contract_db.update_quote(quote['id'], {'expected_version': 2, 'grade': '70/100', 'change_reason': 'Correct grade', 'actor': 'editor'})
    assert revised['version'] == 3 and contract_db.quote_history(quote['id'])[0]['actor'] == 'editor'
    with pytest.raises(QuoteConflict):
        contract_db.update_quote(quote['id'], {'expected_version': 2, 'unit_price': 400, 'change_reason': 'Old version'})
    assert not contract_db.compare_quotes([quote['id'], second['id']])['ranking']
    budget = contract_db.create_procurement({'quote_id': quote['id'], 'quote_version': 3})
    assert budget['total_cost_usd'] == 5500

def test_personal_credentials_live_permissions_and_audited_rollback(contract_db):
    admin = contract_db.create_team_user({'username': 'owner', 'role': 'admin', 'password': 'contract-password-2026'})
    reader = contract_db.audited_write(admin, 'user.created', 'account', lambda: contract_db.create_team_user({'username': 'reader', 'role': 'viewer', 'password': 'contract-password-2026'}))
    session = contract_db.login('reader', 'contract-password-2026', '127.0.0.1')
    token = contract_db.create_user_token(reader['id'], 'Integration')['token']
    assert contract_db.authenticate_session(session['session'])['role'] == 'viewer'
    contract_db.update_team_user(reader['id'], {'version': reader['version'], 'role': 'editor'})
    assert contract_db.authenticate_user_token(token)['role'] == 'editor'
    with pytest.raises(RuntimeError):
        def incomplete():
            contract_db.create_quote(commercial_quote())
            raise RuntimeError('audit transaction must roll back')
        contract_db.audited_write(admin, 'quote.created', 'quote', incomplete)
    assert contract_db.quotes()['total'] == 0
    latest = next(u for u in contract_db.team_users() if u['id'] == reader['id'])
    contract_db.update_team_user(reader['id'], {'version': latest['version'], 'active': False})
    assert contract_db.authenticate_session(session['session']) is None
    assert contract_db.authenticate_user_token(token) is None

def test_durable_quotas_and_failure_backoff_survive_a_second_connection(contract_db):
    restarted = MySQLDatabase(contract_db.settings) if hasattr(contract_db, 'settings') else SQLiteDatabase(contract_db.path)
    now = datetime.now(timezone.utc).timestamp()
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(lambda i: (contract_db if i % 2 else restarted).reserve_provider_request('alpha_vantage', now, 25), range(40)))
    assert sum(row['allowed'] for row in results) == 25
    assert restarted.request_budget('alpha_vantage', now, 25)['used'] == 25
    lease = contract_db.claim_collection('alpha_vantage', now)
    assert lease
    contract_db.finish_collection('alpha_vantage', lease, now=now, next_due=now+60, failures=1, status='error')
    assert restarted.claim_collection('alpha_vantage', now+59, force=True) is None
    assert restarted.claim_collection('alpha_vantage', now+60, force=True)
